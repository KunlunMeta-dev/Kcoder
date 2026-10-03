"""Generate real Office/PDF/image samples and independently reopen valid variants."""
import sys,json,zipfile,datetime
from pathlib import Path
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from openpyxl import Workbook,load_workbook
import xlsxwriter
from PIL import Image,ImageDraw
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.lib.pdfencrypt import StandardEncryption
import msoffcrypto,io
root=Path(sys.argv[1]);root.mkdir(parents=True,exist_ok=True)
cases=[]
def add(name,markers=(),error=None,**more):
 cases.append(dict(name=name,markers=list(markers),error=error,**more))
for name,enc in [('utf8.txt','utf-8'),('utf8-bom.md','utf-8-sig'),('utf16le.txt','utf-16'),('utf16be.md','utf-16-be')]:
 data='WIKI_TEXT_42\n中文资料：价格为42元。\n'.encode(enc)
 if enc=='utf-16-be':data=b'\xfe\xff'+data
 (root/name).write_bytes(data);add(name,['WIKI_TEXT_42','价格为42元'])
doc=Document();doc.add_heading('Wiki Word 测试',0);doc.add_paragraph('WIKI_DOCX_BODY & <42> 中文正文。')
table=doc.add_table(rows=2,cols=2);table.cell(0,0).text='项目';table.cell(0,1).text='数量';table.cell(1,0).text='WIKI_DOCX_TABLE';table.cell(1,1).text='42'
doc.sections[0].header.paragraphs[0].text='HEADER_NOT_PROMISED'
doc.save(root/'word.docx');(root/'unicode.docx').write_bytes((root/'word.docx').read_bytes());add('中文文件.DOCX',['WIKI_DOCX_BODY','WIKI_DOCX_TABLE'],file='unicode.docx');add('word.docx',['WIKI_DOCX_BODY & <42>','WIKI_DOCX_TABLE','42'])
book=Workbook();sheet=book.active;sheet.title='中文表';sheet['A1']='WIKI_XLSX_INLINE';sheet['B1']=42;sheet['C1']=True;sheet['D1']=datetime.date(2026,10,2);sheet['E1']='=B1+8';sheet['A2']='& <中文>'
book.create_sheet('Second')['A1']='SECOND_SHEET_MARKER';book.save(root/'excel-inline.xlsx')
add('excel-inline.xlsx',['Sheet: 中文表','A1: WIKI_XLSX_INLINE','B1: 42','=B1+8 [no cached value]','SECOND_SHEET_MARKER'],observations=['Dates/styles and booleans may remain raw stored values'])
book=xlsxwriter.Workbook(root/'excel-cached.xlsx');sheet=book.add_worksheet('Shared');sheet.write('A1','WIKI_XLSX_SHARED');sheet.write_number('A2',7);sheet.write_number('B2',8);sheet.write_formula('C2','=A2+B2',None,15);sheet.write_url('A3','https://example.com',string='LINK_LABEL');book.close()
add('excel-cached.xlsx',['A1: WIKI_XLSX_SHARED','C2: 15','A3: LINK_LABEL'])
pres=Presentation()
for marker in ['WIKI_PPT_FIRST','WIKI_PPT_SECOND']:
 slide=pres.slides.add_slide(pres.slide_layouts[6]);box=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(7),Inches(2));box.text_frame.text=marker+' 中文幻灯片'
 slide.notes_slide.notes_text_frame.text='NOTES_NOT_PROMISED'
pres.slides._sldIdLst.insert(0,pres.slides._sldIdLst[-1]);pres.save(root/'slides.pptx')
add('slides.pptx',['WIKI_PPT_SECOND','WIKI_PPT_FIRST','中文幻灯片'],ordered=['WIKI_PPT_SECOND','WIKI_PPT_FIRST'])
# UTF-16 XML is valid XML inside OOXML; independent Office readers verify variants.
for source,name,part,verify in [('word.docx','word-utf16-xml.docx','word/document.xml',lambda p: Document(p).paragraphs[1].text),('excel-inline.xlsx','excel-utf16-xml.xlsx','xl/workbook.xml',lambda p: load_workbook(p).active['A1'].value),('slides.pptx','slides-utf16-xml.pptx','ppt/presentation.xml',lambda p: Presentation(p).slides[0].shapes[0].text)]:
 with zipfile.ZipFile(root/source) as old,zipfile.ZipFile(root/name,'w',zipfile.ZIP_DEFLATED) as new:
  for entry in old.infolist():
   data=old.read(entry.filename)
   if entry.filename==part:data=data.decode('utf-8').replace('UTF-8','UTF-16').replace('utf-8','utf-16').encode('utf-16')
   new.writestr(entry,data)
 assert verify(root/name)
 add(name,[next(c['markers'][0] for c in cases if c['name']==source)],edge=True,independentReaderVerified=True)
pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
pdf=canvas.Canvas(str(root/'text.pdf'));pdf.setFont('Helvetica',12);pdf.drawString(40,760,'WIKI_PDF_PAGE1');pdf.showPage();pdf.setFont('STSong-Light',12);pdf.drawString(40,760,'中文资料：价格为42元。');pdf.showPage();pdf.setFont('Helvetica',12);pdf.drawString(40,760,'WIKI_PDF_PAGE3');pdf.save()
add('text.pdf',['WIKI_PDF_PAGE1','中文资料','WIKI_PDF_PAGE3'],pages=[1,2,3])
image=Image.new('RGB',(512,256),'white');draw=ImageDraw.Draw(image);draw.text((25,25),'WIKI_IMAGE_42',fill='black');draw.rectangle((25,65,145,165),fill='blue');draw.text((25,180),'Amount: 42',fill='black')
for name,fmt in [('image.png','PNG'),('image.jpg','JPEG'),('image.jpeg','JPEG'),('image.webp','WEBP')]:
 image.save(root/name,format=fmt);add(name,['NATIVE_IMAGE_INTERPRETATION'],image=True)
pdf=canvas.Canvas(str(root/'scan.pdf'));pdf.drawInlineImage(image,40,500,width=512,height=256);pdf.save();add('scan.pdf',error='needs_vision')
for name in ['old.doc','old.xls','old.ppt','other.gif']:
 (root/name).write_bytes(b'fixture');add(name,error='unsupported_format')
(root/'gbk.txt').write_bytes('中文编码'.encode('gbk'));add('gbk.txt',error='encoding')
(root/'corrupt.docx').write_bytes(b'invalid zip');add('corrupt.docx',error='Invalid or encrypted')
(root/'corrupt.pdf').write_bytes(b'not PDF');add('corrupt.pdf',error='Invalid PDF header')
Image.new('RGB',(16,64)).save(root/'tiny.png');add('tiny.png',error='too small')
Image.new('RGB',(8193,32)).save(root/'wide.png');add('wide.png',error='too large')
(root/'spoof.jpg').write_bytes((root/'image.png').read_bytes());add('spoof.jpg',error='MIME')
(root/'large-text.txt').write_bytes(b'A'*(8*1024*1024+1));add('large-text.txt',error='exceeds 8 MiB')
with open(root/'word.docx','rb') as source,open(root/'encrypted.docx','wb') as dest:
 msoffcrypto.OfficeFile(source).encrypt('fixture-password',dest)
with open(root/'encrypted.docx','rb') as source:
 office=msoffcrypto.OfficeFile(source);assert office.is_encrypted();office.load_key(password='fixture-password');decoded=io.BytesIO();office.decrypt(decoded);assert decoded.getvalue()==(root/'word.docx').read_bytes()
add('encrypted.docx',error='Invalid or encrypted')
pdf=canvas.Canvas(str(root/'encrypted.pdf'),encrypt=StandardEncryption('fixture-password'));pdf.drawString(40,760,'PRIVATE_FIXTURE_TEXT');pdf.save();add('encrypted.pdf',error='encrypted')
(root/'vision-off.png').write_bytes((root/'image.png').read_bytes());add('vision-off.png',error='native vision',noVision=True)
html='<!doctype html><html><head><meta charset="UTF-8"><title>WIKI_HTML_TITLE</title><script>document.write("NEVER_EXECUTE")</script></head><body><h1>中文资料</h1><p>WIKI_HTML_BODY &amp; value</p><ul><li>List evidence</li></ul><table><tr><td>Cell</td><td>42</td></tr></table><div hidden>HIDDEN_HTML</div><img src="https://invalid.example/no-fetch" alt="Diagram label"></body></html>'
(root/'web.html').write_text(html,encoding='utf-8');add('web.html',['WIKI_HTML_TITLE','WIKI_HTML_BODY & value','List evidence','Cell','42','Diagram label'],excluded=['NEVER_EXECUTE','HIDDEN_HTML','invalid.example'])
(root/'web.htm').write_bytes(html.encode('utf-16'));add('web.htm',['WIKI_HTML_TITLE','中文资料'])
(root/'web-gbk.html').write_bytes('<meta charset=gbk><p>WIKI_HTML_GBK 中文资料</p>'.encode('gbk'));add('web-gbk.html',['WIKI_HTML_GBK 中文资料'])
(root/'empty.html').write_text('<html><script>only code</script></html>');add('empty.html',error='no extractable text')
(root/'manifest.json').write_text(json.dumps(cases,ensure_ascii=False,indent=2))
print('Generated',len(cases),'real and negative format samples')
